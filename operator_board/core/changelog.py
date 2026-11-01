"""A record of every setting changed from Operator Board, tied to the slate, take, timecode, date and
time it was changed at, for operators, DITs, data wranglers and VFX houses.

Each change is a ChangeRecord: what was changed (category, item, setting), the new and the previous
value, and where it stood on set at that moment. The log is kept in ``change_log.jsonl`` (one record per
line, so it survives a crash) and exports to CSV or to a spreadsheet (.xlsx).

Changes that can be put back (Remote Objects) also carry ``kind`` and ``ref``: what to call to apply a
value again, which Undo uses.
"""
from __future__ import annotations

import csv
import datetime
import json
import logging
import os
import threading
from dataclasses import asdict, dataclass, field
from typing import Any, Callable, Dict, List, Optional, Sequence

from operator_board.core.config import app_home

log = logging.getLogger(__name__)

COLUMNS = [("date", "Date"), ("time", "Time"), ("timecode", "Timecode"), ("slate", "Slate"), ("take", "Take"),
           ("category", "Category"), ("item", "Item"), ("setting", "Setting"), ("value", "Value"),
           ("previous", "Previous value"), ("source", "Source"), ("result", "Result"), ("project", "Project"),
           ("level", "Level")]


def show(value: Any) -> str:
    """A value as readable text for exports (lists and dicts as compact JSON)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "on" if value else "off"
    if isinstance(value, float):
        return f"{value:.6g}"
    if isinstance(value, (list, dict)):
        return json.dumps(value, separators=(", ", ": "))
    return str(value)


@dataclass
class ChangeRecord:
    date: str = ""
    time: str = ""
    timecode: str = ""
    slate: str = ""
    take: int = 0
    category: str = ""
    item: str = ""
    setting: str = ""
    value: Any = None
    previous: Any = None
    source: str = "Operator Board"
    result: str = "ok"
    project: str = ""
    level: str = ""
    kind: str = ""  # how to apply the value again (Undo, presets): "quick" | "property" | "preset" | ""
    ref: Dict[str, Any] = field(default_factory=dict)

    @property
    def key(self) -> tuple:
        """Which setting this is (changes to the same setting share it)."""
        return (self.kind, json.dumps(self.ref, sort_keys=True), self.setting)


class ChangeLog:
    def __init__(self, context: Callable[[], Dict[str, Any]], path: Optional[str] = None) -> None:
        """``context()`` returns slate, take, timecode, project and level at the moment of a change."""
        self.context = context
        self.path = path or os.path.join(app_home(), "change_log.jsonl")
        self.records: List[ChangeRecord] = []
        self._lock = threading.Lock()
        self.listeners: List[Callable[[ChangeRecord], None]] = []
        self._load()

    def _load(self) -> None:
        if not os.path.exists(self.path):
            return
        try:
            with open(self.path, "r", encoding="utf-8") as fh:
                for line in fh:
                    line = line.strip()
                    if line:
                        data = json.loads(line)
                        self.records.append(ChangeRecord(**{k: v for k, v in data.items()
                                                            if k in ChangeRecord.__dataclass_fields__}))
        except (OSError, ValueError, TypeError) as exc:
            log.warning("Could not read the change log %s: %s", self.path, exc)

    def record(self, category: str, item: str, setting: str, value: Any, previous: Any = None,
               source: str = "Operator Board", result: str = "ok", kind: str = "",
               ref: Optional[Dict[str, Any]] = None) -> ChangeRecord:
        now = datetime.datetime.now()
        try:
            ctx = self.context() or {}
        except Exception:  # noqa: BLE001 - a change is still recorded without its context
            ctx = {}
        rec = ChangeRecord(date=now.strftime("%Y-%m-%d"), time=now.strftime("%H:%M:%S"),
                           timecode=ctx.get("timecode") or "", slate=ctx.get("slate", ""), take=ctx.get("take", 0),
                           category=category, item=item, setting=setting, value=value, previous=previous,
                           source=source, result=result, project=ctx.get("project", ""), level=ctx.get("level", ""),
                           kind=kind, ref=dict(ref or {}))
        with self._lock:
            self.records.append(rec)
            try:
                os.makedirs(os.path.dirname(os.path.abspath(self.path)), exist_ok=True)
                with open(self.path, "a", encoding="utf-8") as fh:
                    fh.write(json.dumps(asdict(rec), default=str) + "\n")
            except OSError as exc:
                log.warning("Could not write the change log: %s", exc)
        for listener in list(self.listeners):
            try:
                listener(rec)
            except Exception:  # noqa: BLE001
                log.exception("Change log listener failed")
        return rec

    def clear(self) -> None:
        with self._lock:
            self.records = []
            try:
                if os.path.exists(self.path):
                    os.replace(self.path, self.path + ".old")  # keep the previous log as a backup
            except OSError as exc:
                log.warning("Could not clear the change log: %s", exc)

    # ---- exports ------------------------------------------------------------------------
    def export_csv(self, path: str, records: Optional[Sequence[ChangeRecord]] = None) -> int:
        rows = list(self.records if records is None else records)
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:  # BOM: Excel opens it as UTF-8
            w = csv.writer(fh)
            w.writerow([title for _key, title in COLUMNS])
            for r in rows:
                w.writerow([show(getattr(r, key)) for key, _title in COLUMNS])
        return len(rows)


def export_workbook(path: str, sheets: Sequence[tuple]) -> None:
    """Write an .xlsx with one sheet per (title, headers, rows): bold frozen header row, filters and
    column widths sized to the content, so it reads well in Excel, Numbers or Google Sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        ws.append(list(headers))
        for row in rows:
            ws.append([show(v) if isinstance(v, (list, dict, bool)) or v is None else v for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for col, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(r[col - 1])) for r in rows if col - 1 < len(r)])
            ws.column_dimensions[get_column_letter(col)].width = min(60, max(10, longest + 2))
    wb.save(path)
